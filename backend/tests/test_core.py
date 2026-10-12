import io
import re

import pytest
from docx import Document
from pypdf import PdfReader

from app.models.schemas import Experience, JobInput, Profile
from app.seed import initial_profile
from app.services import document_exporter, llm_provider, pipeline
from app.services.catalog import CATALOG, find_terms
from app.services.job_matcher import analyze, build_evidence_index
from app.services.job_parser import parse_job
from app.services.resume_generator import generate_resume


# ------------------------------------------------------------ parsing
def test_parser_extracts_title_company_levels(job_text):
    job = parse_job(job_text)
    assert job.title == "Analista DevOps Pleno" and job.company == "Acme Tecnologia"
    assert job.seniority == "Pleno"
    assert job.years_required == 3
    assert job.english_required and job.degree_required
    for t in ("Docker", "Kubernetes", "AWS", "Terraform", "GitHub Actions"):
        assert t in job.technologies_required
    assert "Kafka" in job.technologies_desired
    assert "Kafka" not in job.technologies_required


def test_parser_user_overrides_and_word_boundaries():
    job = parse_job("Desenvolvedor JavaScript com Java e .NET. Conhecimento de Go é irrelevante, usamos golang.", "Acme", "Dev")
    assert (job.company, job.title) == ("Acme", "Dev")
    kw = job.keywords
    assert "JavaScript" in kw and "Java" in kw and ".NET" in kw and "Go" in kw
    assert "Java" in find_terms("Java") and "Java" not in find_terms("JavaScript")


def test_job_input_validation():
    with pytest.raises(ValueError):
        JobInput(job_description="curta")
    with pytest.raises(ValueError):
        JobInput(job_description="x" * 40_000)


# ------------------------------------------------------------ perfil
def test_profile_integrity_no_invented_data():
    p = initial_profile()
    assert [x.company for x in p.experiences] == ["Produto Reativo"] * 3 and p.languages == []
    assert (p.personal.email, p.personal.phone, p.personal.location) == ("christian@milfont.org", "+55 85 98191-3769", "")
    assert {c.name for c in p.certifications} == {
        "Google Cloud Generative AI Leader", "Microsoft Azure AZ-104", "AWS re/Start", "Dynatrace Essentials",
    }
    assert p.personal.linkedin.endswith("christianmilfont")
    blob = p.model_dump_json()
    assert not re.search(r"\d+\s*%|\d+\s*anos", blob)
    assert Profile.model_validate_json(blob) == p


# ------------------------------------------------------------ score
def test_score_is_explainable_and_honest(job_text):
    p = initial_profile()
    a = analyze(parse_job(job_text), p)
    assert 0 <= a.score <= 100
    assert abs(sum(c.points for c in a.score_components) / sum(c.max_points for c in a.score_components) * 100 - a.score) <= 1
    matched = {m.requirement for m in a.match}
    assert {"Docker", "Kubernetes", "AWS", "Terraform"} <= matched
    gaps = {g.requirement: g for g in a.gaps}
    assert "Prometheus" in gaps and "relacionado" in gaps["Prometheus"].evidence.lower()  # Grafana é relacionado
    assert "Kafka" in gaps
    risks = " ".join(a.risks)
    assert "3+ anos" not in risks and "inglês" in risks.lower()
    assert "ATS" in a.summary  # não promete aprovação


def test_score_rewards_better_fit():
    p = initial_profile()
    strong = analyze(parse_job("Estágio DevOps. Requisitos: Docker, Kubernetes, AWS, Terraform, CI/CD, GitHub Actions."), p)
    weak = analyze(parse_job("Desenvolvedor Pleno. Requisitos: PHP, Laravel, Ruby, Angular, Oracle. 5 anos de experiência."), p)
    assert strong.score > weak.score + 30


def test_experience_in_profile_removes_risk(job_text):
    p = initial_profile()
    p.experiences = [Experience(company="X", role="Dev", start="2020", end="2022")]
    a = analyze(parse_job(job_text), p)
    assert not any("experiência profissional cadastrada" in r for r in a.risks)


# ------------------------------------------------------------ currículo
def test_resume_only_uses_supported_terms(job_text):
    p = initial_profile()
    a = pipeline.analyze_job(JobInput(job_description=job_text), p)
    resume, notices = pipeline.build_resume(a, p)
    idx = set(build_evidence_index(p))
    text = resume.model_dump_json()
    # nenhuma tecnologia do catálogo que não esteja no perfil (Kafka, Prometheus aparecem na vaga mas não no perfil)
    for term in find_terms(text):
        assert term in idx, f"{term} sem respaldo no perfil"
    assert "Kafka" not in text and "Prometheus" not in text
    assert len(resume.experience) == 3 and resume.experience[0].period == "2024 – 2025 | Remoto" and resume.languages == []
    assert not re.search(r"\d+\s*\+?\s*anos", resume.summary)
    assert any("sem IA" in n for n in notices)
    assert a.mode == "rules"
    assert resume.name == p.personal.full_name
    assert [l.label for l in resume.links] == ["LinkedIn", "GitHub", "Portfólio"]


def test_resume_prioritizes_and_filters(job_text):
    p = initial_profile()
    r, _ = generate_resume(parse_job("Dev .NET\nRequisitos: C#, .NET, MySQL, JWT, APIs REST"), None, p)
    assert r.projects[0].name == "SkillBridge"
    assert r.skills[0].category == "Desenvolvimento"
    assert len(r.projects) <= 3


# ------------------------------------------------------------ LLM
def test_llm_disabled_by_default(monkeypatch):
    monkeypatch.delenv("LLM_PROVIDER", raising=False)
    s = llm_provider.get_status()
    assert not s.ai_enabled and s.mode == "rules" and "sem IA" in s.message


def test_llm_validation_rejects_invention():
    allowed = {"Docker", "AWS"}
    assert llm_provider.validate_generated_text("Profissional com 5 anos de experiência", allowed, False)
    assert llm_provider.validate_generated_text("Conhecimento em Kafka", allowed, False)
    assert llm_provider.validate_generated_text("Atuou como DevOps", allowed, False)
    assert llm_provider.validate_generated_text("Conhecimentos em Docker e AWS", allowed, False) is None


def test_llm_failure_falls_back_with_notice(monkeypatch, job_text):
    monkeypatch.setenv("LLM_PROVIDER", "openai")
    monkeypatch.setenv("LLM_MODEL", "x")
    monkeypatch.setenv("LLM_API_KEY", "k")

    def boom(*a, **k):
        raise llm_provider.LLMError("Falha ao chamar o provedor de IA (Teste).")

    monkeypatch.setattr(llm_provider, "complete_json", boom)
    p = initial_profile()
    a = pipeline.analyze_job(JobInput(job_description=job_text), p)
    assert a.mode == "rules" and any("Falha" in n for n in a.notices)
    resume, notices = pipeline.build_resume(a, p)
    assert resume.summary and any("Falha" in n for n in notices)


def test_llm_success_path(monkeypatch, job_text):
    monkeypatch.setenv("LLM_PROVIDER", "openai")
    monkeypatch.setenv("LLM_MODEL", "x")
    monkeypatch.setenv("LLM_API_KEY", "k")

    def fake(system, user, cfg=None):
        if "extraem" in system or "extrai" in system:
            return {"title": "Analista DevOps", "required": ["Docker"], "technologies_required": ["Docker", "Inventado123"]}
        return {"title": "Analista DevOps", "summary": "Formação em ADS e pós-graduação em DevOps, com projetos usando Docker e AWS."}

    monkeypatch.setattr(llm_provider, "complete_json", fake)
    p = initial_profile()
    a = pipeline.analyze_job(JobInput(job_description=job_text), p)
    assert a.mode == "llm"
    assert "Inventado123" not in a.job.technologies_required  # termo ausente da vaga é descartado
    resume, _ = pipeline.build_resume(a, p)
    assert resume.summary.startswith("Formação em ADS")


# ------------------------------------------------------------ exportação
@pytest.fixture()
def resume(job_text):
    p = initial_profile()
    a = pipeline.analyze_job(JobInput(job_description=job_text), p)
    return pipeline.build_resume(a, p)[0]


def test_pdf_export_selectable_text_links_and_pages(resume):
    data = document_exporter.export_pdf(resume)
    assert data.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(data))
    assert 1 <= len(reader.pages) <= 2
    text = "\n".join(pg.extract_text() for pg in reader.pages)
    assert "Christian Milfont" in text and "Competências técnicas".upper() in text.upper()
    assert "Formação acadêmica".upper() in text.upper()
    uris = []
    for pg in reader.pages:
        for a in pg.get("/Annots") or []:
            uri = a.get_object().get("/A", {}).get("/URI")
            if uri:
                uris.append(str(uri))
    assert "https://github.com/christianmilfont" in uris


def test_pdf_handles_special_chars_and_empty():
    from app.models.schemas import ResumeContent

    r = ResumeContent(name="Ana <b>&</b> Ç", title="Dev → Ops", summary="Texto com “aspas” & <tags>")
    assert document_exporter.export_pdf(r).startswith(b"%PDF")
    assert document_exporter.export_docx(r)[:2] == b"PK"


def test_docx_export_content_and_hyperlinks(resume):
    data = document_exporter.export_docx(resume)
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Christian Milfont" in text and "COMPETÊNCIAS TÉCNICAS" in text
    rels = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://linkedin.com/in/christianmilfont" in rels


def test_all_sections_order_respected(resume):
    resume.section_order = ["education", "summary"]
    text = "\n".join(p.text for p in Document(io.BytesIO(document_exporter.export_docx(resume))).paragraphs)
    assert text.index("FORMAÇÃO") < text.index("RESUMO")
    assert "COMPETÊNCIAS" not in text


def test_title_with_location_suffix_is_not_company():
    job = parse_job("Engenheiro de Software - Remoto\nRequisitos: Python, SQL e APIs REST para o time de produto.")
    assert job.title == "Engenheiro de Software - Remoto" and job.company == ""


def test_title_with_state_suffix_is_not_company():
    job = parse_job("Analista de Dados - SP\nRequisitos: SQL, Python e Power BI para o time de dados.")
    assert job.company == ""


def test_no_experience_profile_raises_years_and_seniority_risks(job_text):
    p = initial_profile()
    p.experiences = []
    risks = " ".join(analyze(parse_job(job_text), p).risks)
    assert "3+ anos" in risks and "Pleno" in risks


def test_resume_is_concise_and_fits_one_page(resume):
    assert all(len(e.bullets) <= 3 for e in resume.experience)
    assert len(resume.projects) <= 2
    assert len(PdfReader(io.BytesIO(document_exporter.export_pdf(resume))).pages) == 1
    _, pages = document_exporter._build_pdf(resume, 10)
    assert pages == 1  # cabe sem precisar reduzir a fonte abaixo de 10pt
