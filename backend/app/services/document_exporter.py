"""Exportação do currículo para PDF (reportlab, texto selecionável) e DOCX (python-docx). Layout de coluna única."""
from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.models.schemas import ResumeContent

SECTION_TITLES = {
    "summary": "Resumo profissional", "skills": "Competências técnicas", "experience": "Experiência profissional",
    "projects": "Projetos", "education": "Formação acadêmica", "certifications": "Certificações",
    "achievements": "Conquistas", "languages": "Idiomas",
}
ACCENT = colors.HexColor("#1f3a5f")
_CP1252_FIX = {"→": "->", "’": "'", "“": '"', "”": '"', " ": " ", "​": ""}


def _pdf_safe(text: str) -> str:
    """Fontes base do PDF cobrem Latin-1/cp1252; troca o que ficaria ilegível."""
    for k, v in _CP1252_FIX.items():
        text = text.replace(k, v)
    return text.encode("cp1252", errors="replace").decode("cp1252")


def _sections(r: ResumeContent) -> list[str]:
    seen, order = set(), []
    for s in r.section_order:
        if s not in seen:
            seen.add(s)
            order.append(s)
    return order


def _has_content(r: ResumeContent, s: str) -> bool:
    return bool({
        "summary": r.summary.strip(), "skills": r.skills, "experience": r.experience, "projects": r.projects,
        "education": r.education, "certifications": r.certifications, "achievements": r.achievements,
        "languages": r.languages,
    }[s])


def safe_filename(r: ResumeContent, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", f"Curriculo_{r.name}_{r.title}".encode("ascii", "ignore").decode()).strip("_")
    return f"{base[:80] or 'Curriculo'}.{ext}"


# ============================================================ PDF
def _pdf_styles(size: float) -> dict[str, ParagraphStyle]:
    base = dict(fontName="Helvetica", fontSize=size, leading=size * 1.26, textColor=colors.black)
    return {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=size + 9, leading=size + 12, alignment=TA_CENTER, textColor=ACCENT),
        "title": ParagraphStyle("title", fontName="Helvetica", fontSize=size + 2, leading=size + 5, alignment=TA_CENTER),
        "contact": ParagraphStyle("contact", alignment=TA_CENTER, fontSize=size - 0.5, leading=(size - 0.5) * 1.35, fontName="Helvetica"),
        "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=size + 1, leading=size + 4, textColor=ACCENT, spaceBefore=size * 0.7, spaceAfter=1),
        "p": ParagraphStyle("p", **base),
        "item": ParagraphStyle("item", fontName="Helvetica-Bold", fontSize=size, leading=size * 1.26, spaceBefore=2),
        "meta": ParagraphStyle("meta", fontName="Helvetica-Oblique", fontSize=size - 0.5, leading=size * 1.25, textColor=colors.HexColor("#444444")),
    }


def _link(url: str, label: str) -> str:
    return f'<a href="{escape(url, {chr(34): "&quot;"})}" color="#1f3a5f">{escape(label)}</a>'


def _build_pdf(r: ResumeContent, size: float) -> tuple[bytes, int]:
    st = _pdf_styles(size)
    t = lambda s: escape(_pdf_safe(s))  # noqa: E731
    story: list = []
    if r.name:
        story.append(Paragraph(t(r.name), st["name"]))
    if r.title:
        story.append(Paragraph(t(r.title), st["title"]))
    contact_bits = [t(c) for c in r.contacts] + [_link(l.url, _pdf_safe(l.label) + ": " + _pdf_safe(l.url.replace("https://", "").replace("http://", "").rstrip("/"))) for l in r.links]
    if contact_bits:
        story.append(Paragraph(" | ".join(contact_bits), st["contact"]))

    def bullets(items: list[str]) -> ListFlowable:
        return ListFlowable(
            [ListItem(Paragraph(t(b), st["p"]), leftIndent=12) for b in items if b.strip()],
            bulletType="bullet", start="•", leftIndent=12, bulletFontSize=size - 1,
        )

    for sec in _sections(r):
        if not _has_content(r, sec):
            continue
        story.append(Paragraph(SECTION_TITLES[sec].upper(), st["h"]))
        story.append(HRFlowable(width="100%", thickness=0.6, color=ACCENT, spaceAfter=3))
        if sec == "summary":
            story.append(Paragraph(t(r.summary), st["p"]))
        elif sec == "skills":
            for g in r.skills:
                story.append(Paragraph(f"<b>{t(g.category)}:</b> {t(', '.join(g.items))}", st["p"]))
        elif sec == "experience":
            for e in r.experience:
                story.append(Paragraph(f"{t(e.role)} — {t(e.company)}" + (f" ({t(e.period)})" if e.period else ""), st["item"]))
                if e.bullets:
                    story.append(bullets(e.bullets))
        elif sec == "projects":
            for p in r.projects:
                head = t(p.name) + (f" — {t(p.subtitle)}" if p.subtitle else "")
                story.append(Paragraph(head, st["item"]))
                if p.description:
                    story.append(Paragraph(t(p.description), st["p"]))
                if p.bullets:
                    story.append(bullets(p.bullets))
                if p.technologies:
                    story.append(Paragraph(f"<i>Tecnologias:</i> {t(', '.join(p.technologies))}", st["meta"]))
                if p.link:
                    story.append(Paragraph(_link(p.link, _pdf_safe(p.link)), st["meta"]))
        elif sec == "education":
            for e in r.education:
                story.append(Paragraph(f"{t(e.course)}" + (f" — {t(e.institution)}" if e.institution else ""), st["item"]))
                meta = " | ".join(x for x in (e.period, e.notes) if x)
                if meta:
                    story.append(Paragraph(t(meta), st["meta"]))
        elif sec in ("certifications", "achievements", "languages"):
            story.append(bullets(getattr(r, sec)))
        story.append(Spacer(1, 1))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.2 * cm, bottomMargin=1.1 * cm,
        title=_pdf_safe(f"Currículo - {r.name}"), author=_pdf_safe(r.name),
    )
    doc.build(story)
    return buf.getvalue(), doc.page


def export_pdf(r: ResumeContent) -> bytes:
    """Tenta caber em uma página reduzindo a fonte até 9pt; senão usa duas páginas (ou mais, se o texto exigir)."""
    result = b""
    for size in (10.5, 10, 9.5, 9):
        result, pages = _build_pdf(r, size)
        if pages <= 1:
            return result
    return _build_pdf(r, 10)[0]


# ============================================================ DOCX
def _add_hyperlink(par, url: str, text: str, size: float) -> None:
    rid = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F3A5F")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    for el in (color, u, sz):
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def _bottom_border(par) -> None:
    ppr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
        b.set(qn(k), v)
    borders.append(b)
    ppr.append(borders)


def export_docx(r: ResumeContent) -> bytes:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    d.core_properties.title = f"Currículo - {r.name}"
    d.core_properties.author = r.name

    def para(text: str = "", bold=False, italic=False, size=10.5, align=None, color=None, before=0, after=0, style=None):
        p = d.add_paragraph(style=style)
        if text:
            run = p.add_run(text)
            run.bold, run.italic = bold, italic
            run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        if align:
            p.alignment = align
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        return p

    def bullet_list(items: list[str]) -> None:
        for b in items:
            if b.strip():
                p = para(b, style="List Bullet")
                p.paragraph_format.left_indent = Cm(0.6)

    center = WD_ALIGN_PARAGRAPH.CENTER
    if r.name:
        para(r.name, bold=True, size=20, align=center, color="1F3A5F")
    if r.title:
        para(r.title, size=12, align=center, after=2)
    if r.contacts or r.links:
        p = para(align=center, after=2)
        first = True
        for c in r.contacts:
            run = p.add_run(("" if first else " | ") + c)
            run.font.size = Pt(10)
            first = False
        for l in r.links:
            if not first:
                run = p.add_run(" | ")
                run.font.size = Pt(10)
            _add_hyperlink(p, l.url, f"{l.label}: {l.url.replace('https://', '').replace('http://', '').rstrip('/')}", 10)
            first = False

    for s in _sections(r):
        if not _has_content(r, s):
            continue
        h = para(SECTION_TITLES[s].upper(), bold=True, size=11, color="1F3A5F", before=8, after=3)
        _bottom_border(h)
        if s == "summary":
            para(r.summary)
        elif s == "skills":
            for g in r.skills:
                p = d.add_paragraph()
                p.add_run(f"{g.category}: ").bold = True
                p.add_run(", ".join(g.items))
        elif s == "experience":
            for e in r.experience:
                para(f"{e.role} — {e.company}" + (f" ({e.period})" if e.period else ""), bold=True, before=3)
                bullet_list(e.bullets)
        elif s == "projects":
            for pr in r.projects:
                para(pr.name + (f" — {pr.subtitle}" if pr.subtitle else ""), bold=True, before=3)
                if pr.description:
                    para(pr.description)
                bullet_list(pr.bullets)
                if pr.technologies:
                    p = d.add_paragraph()
                    p.add_run("Tecnologias: ").italic = True
                    p.add_run(", ".join(pr.technologies)).italic = True
                if pr.link:
                    _add_hyperlink(d.add_paragraph(), pr.link, pr.link, 10)
        elif s == "education":
            for e in r.education:
                para(e.course + (f" — {e.institution}" if e.institution else ""), bold=True, before=3)
                meta = " | ".join(x for x in (e.period, e.notes) if x)
                if meta:
                    para(meta, italic=True, size=10)
        else:
            bullet_list(getattr(r, s))

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
